"""Builds docs/LoanSeva_Build_and_Setup_Guide.docx. Run: python3 scripts/gen_guide.py"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x00, 0x38, 0x74)
RED = RGBColor(0xED, 0x1C, 0x24)
GREY = RGBColor(0x5F, 0x63, 0x68)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2)
base = doc.styles['Normal']
base.font.name = 'Calibri'
base.font.size = Pt(10.5)
for name, size, color in (('Heading 1', 16, NAVY), ('Heading 2', 12.5, NAVY), ('Heading 3', 11, RED)):
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = True


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hex_fill)
    tcPr.append(sh)


def table(rows, widths=None, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ''
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(9.5)
            if header and r == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(cell, '003874')
            elif c == 0:
                run.font.bold = True
            if widths:
                cell.width = Cm(widths[c])
    doc.add_paragraph()
    return t


def para(txt, bold_lead=None, style=None):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(txt)
    return p


def bullets(items):
    for i in items:
        if isinstance(i, tuple):
            para(i[1], bold_lead=i[0] + ' ', style='List Bullet')
        else:
            doc.add_paragraph(i, style='List Bullet')


def steps(items):
    for i in items:
        if isinstance(i, tuple):
            para(i[1], bold_lead=i[0] + ' ', style='List Number')
        else:
            doc.add_paragraph(i, style='List Number')


def code(txt):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(txt)
    r.font.name = 'Menlo'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)
    pPr = p._p.get_or_add_pPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:fill'), 'F3F4F6')
    pPr.append(sh)


def note(txt):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run('Note: ')
    r.bold = True
    r.font.color.rgb = RED
    p.add_run(txt).italic = True


# ------------------------------------------------------------------ Cover
p = doc.add_paragraph()
r = p.add_run('KOTAK MAHINDRA BANK × SALESFORCE')
r.bold = True
r.font.color.rgb = RED
r.font.size = Pt(11)
p = doc.add_paragraph()
r = p.add_run('LoanSeva by Kotak: WhatsApp Banking on Salesforce')
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = NAVY
p = doc.add_paragraph()
r = p.add_run('What we built in Salesforce, and how to enable it')
r.font.size = Pt(14)
r.font.color.rgb = GREY
table([
    ('Org', 'kotak-demo (Salesforce trial org 00DgL00000eqlv3UAA), source API 65.0'),
    ('Source', 'Git repository kotak-loanseva (force-app/main/default); GitHub: tilotianikita-coder/Whatsapp-Banking'),
    ('Status', 'Proof of concept, deployed and working end to end; 118 Apex tests pass'),
    ('Data', 'All rates, fees, CIBIL, KYC and bank-statement results are illustrative demo values. No live bank systems are called.'),
], widths=[3, 13.5], header=False)

# ------------------------------------------------------------------ 1
doc.add_heading('1. Summary', level=1)
para('LoanSeva turns WhatsApp into a full lending channel for Kotak personal loans. An existing customer goes from a QR code or a "Hi" '
     'to a verified, personalised Top-Up offer, a reference number, an RM callback and a sanctioned loan account in one WhatsApp thread. '
     'A customer without a pre-approved offer can apply for a new personal loan, share documents as photos and get an underwriting decision from an RM. '
     'Everything the customer says is written to Salesforce as it happens, so the RM sees the same journey, the bank keeps one record, and every drop-off becomes a warm lead.')
para('The design separates the conversation from the decisions. The Agentforce agent understands the customer and relays replies; '
     'a deterministic Apex engine owns the journey state, validation, EMI maths, eligibility and copy in three languages. '
     'Numbers never come from the language model, so the same inputs always produce the same answer, and the whole journey is unit-tested.',
     bold_lead='Design principle. ')
para('Run the loan journey natively on Salesforce (Approach 2): a dedicated Kotak Loans WhatsApp number on Salesforce native WhatsApp, with the Agentforce agent, '
     'the journey engine and the LOS on the same platform. It is the least to build (no BSP connector or hand-over logic) and one vendor for the loan journey. '
     'The main number stays on Gupshup and links across; later it can move to Salesforce too. Approach 1 (Gupshup in front, Salesforce via BYOC) keeps one number '
     'but needs Gupshup to build the connector and the hand-over; section 2.6 shows that path, which is also how the PoC is connected.',
     bold_lead='Recommended approach. ')

# ------------------------------------------------------------------ 2
doc.add_heading('2. What we built', level=1)
doc.add_heading('2.1 Customer journeys on WhatsApp', level=2)
table([
    ('Journey', 'What the customer experiences', 'What happens in Salesforce'),
    ('Entry (QR / link)', 'Scans a QR or taps a link, enters the mobile number, ticks consent, and gets a WhatsApp greeting', 'Public Salesforce Site page publishes a platform event; the subscriber matches the Lead/Contact, creates a Loan Journey and sends the WhatsApp message'),
    ('Personal Loan Top-Up', 'Verifies with name, mobile and OTP; sees the masked existing loan, EMI and repayment track; answers 4 questions; gets a priced offer with editable tenure; accepts and picks a callback slot', 'Offer priced from Existing_Loan__c (pre-approved limit, on-time EMIs, FOIR). Loan Application created with a KMB-PLTU reference. Pre-approved Top-Ups skip underwriting, the LAN is created as Sanctioned, and the application goes to the RM Callback queue with a Task and Event'),
    ('New personal loan', '10 questions (prefilled for existing customers), indicative offer, photos of PAN, bank statement and salary slips (each can be skipped), review and consent', 'Application with KMB-PL reference, documents filed as Loan Documents, mock underwriting runs six checks and sets a recommendation, RM decides in the LOS app'),
    ('Shortcuts at any step', '"EMI 3 lakh 36 months" shows an EMI table; "Status" shows where the application is; FAQs answered from approved content', 'Answer ends with the pending question, so the journey resumes where it was'),
    ('Languages and voice', 'English, Hindi or Marathi, switch any time; answer any step with a voice note', 'All copy in LoanSeva_Message__mdt (EN/HI/MR). Voice notes are transcribed by ElevenLabs speech-to-text; PAN shown masked, OTP never echoed'),
    ('Talk to a specialist', 'Types "agent", "RM" or picks menu option 3 (also in Hindi/Marathi)', 'RM Task created; the agent escalates and Omni-Channel routes the chat, with transcript, to the human queue'),
    ('Drop-off and nudges', 'If the customer goes quiet, they get up to two reminder messages', 'Scheduled job marks the drop point, creates an RM Task and sends nudge 1 and 2 (demo: 2 and 5 minutes; production: 24 h / 72 h)'),
    ('Feedback', 'Rates the experience at the end', 'Journey_Feedback__c record'),
], widths=[3, 6.2, 7.3])

doc.add_heading('2.2 Loan Origination (LoanSeva LOS app)', level=2)
bullets([
    ('Credit Review (top of the Loan Application page).', 'Horizontal stage gates: Application → Documents → KYC (NSDL) → Bureau (CIBIL) → Banking (Perfios) → Employment → Policy (FOIR) → Fraud (Hunter) → RM decision → Loan account. Green pass, yellow refer/pending, red fail, blue awaiting RM, grey not reached. Clicking a gate opens the mock provider report.'),
    ('Pre-approved Top-Up view.', 'Not-required steps (Documents, Bureau & checks, Credit decision) are hatched; Application, Identity, Pre-approval and Offer are green; RM callback is blue ("RM queue"); Loan account is green with the masked LAN and "Sanctioned".'),
    ('RM credit summary.', 'Verdict, narrative, strengths, risks and conditions before disbursal, generated by rules (LoanSevaCreditMemo). For Top-Ups: "Pre-approved Top-Up - no fresh underwriting needed".'),
    ('Underwriting Scorecard (sidebar).', 'Approve & create LAN, Reject, Re-run underwriting; decision dialog requires notes when going against the recommendation. Documents list with RM upload, Verified/Rejected marking. Top-Ups show the pre-approved terms and the sanctioned LAN instead.'),
    ('Loan account.', 'Approval or pre-approved Top-Up creates a Sanctioned Existing_Loan__c with LAN PL + yyMMdd + 8-digit sequence; the customer gets only the masked LAN on WhatsApp.'),
    ('RM Callback queue.', 'Queue LoanSeva_RM_Callbacks owns accepted Top-Ups. App LoanSeva LOS → Loan Applications → list view "RM Callback Queue". Also list views All Applications, All Journeys, All Loans.'),
    ('Currency.', 'Multi-currency is enabled with INR as the corporate currency; triggers stamp INR on new loans, journeys and applications.'),
])

doc.add_heading('2.3 Mock underwriting rules (demo PANs)', level=2)
para('PAN format AAAPS9999X: 4th character P (individual), 5th character the surname initial. The last character picks the bureau outcome; results are seeded from the PAN, so a file always shows the same data.')
table([
    ('PAN example (for "Asha Patil")', 'CIBIL', 'Recommendation'),
    ('ABCPP1234K + bank statement + Category A employer', '777', 'Approve'),
    ('ABCPP1234K, bank statement skipped', '777', 'Refer (statement pending)'),
    ('ABCPP1234R (ends in R)', '695', 'Refer'),
    ('ABCPP1234Z (ends in Z)', '630', 'Decline'),
    ('ABCPS1234K (surname initial mismatch)', '777', 'Refer (name check)'),
], widths=[8, 2.5, 6])
para('Self-employed applicants and Category B employers without salary slips are always Refer; a disposable email domain gives a fraud Refer; the same PAN on a different mobile fails dedupe.')

doc.add_heading('2.4 Architecture', level=2)
table([
    ('Layer', 'Components'),
    ('Channel', 'WhatsApp_BYOC messaging channel (Bring Your Own Channel, from the installed "WhatsApp BYOC Demo (Meta)" package wpbyocdemo3, unmodified). In production the recommended channel is native WhatsApp in Salesforce Messaging (Approach 2); Kotak\'s Gupshup number through BYOC is the alternative (Approach 1).'),
    ('Routing', 'Omni-Flow LoanSeva_WhatsApp_Inbound_Routing sends WhatsApp sessions to the agent; fallback queue SDO_Service_Messaging for human hand-off.'),
    ('Agent', 'Agentforce Service Agent "LoanSeva by Kotak" (classic metadata: Bot, planner bundle, 2 topics, invocable actions). Each action is a thin wrapper over the engine.'),
    ('Engine', 'LoanSevaEngine (state machine), LoanSevaEligibilityEngine, LoanSevaNewLoan, LoanSevaShortcuts, LoanSevaFaqMatcher, LoanSevaI18n, LoanSevaVoice, LoanSevaJourneyService.'),
    ('LOS', 'LoanSevaUnderwriting (mock checks), LoanSevaReport (provider reports), LoanSevaCreditMemo, LoanSevaDecisionService (approve/reject/sanction), LoanSevaDocumentService, LWCs loanSevaCreditReview, loanSevaScorecard, loanSevaDecisionModal.'),
    ('Outbound', 'LoanSevaWhatsAppSender (Queueable, sendConversationMessages via the LoanSeva_Text messaging component) and LoanSevaDropOffJob (scheduled).'),
    ('Entry', 'Salesforce Site "LoanSeva" with LoanSevaRegister Visualforce page; LoanSeva_Event__e platform event; QR in docs/.'),
    ('Integrations', 'ElevenLabs speech-to-text via Named Credential / External Credential LoanSeva_ElevenLabs. Bureau, NSDL, Perfios and core banking are mocked; production would call them through MuleSoft or Named Credentials.'),
], widths=[3, 13.5])

doc.add_heading('2.5 Component inventory', level=2)
table([
    ('Type', 'Count', 'Items'),
    ('Custom objects', '6', 'Loan_Journey__c, Existing_Loan__c, Loan_Application__c, Loan_Document__c, Underwriting_Check__c, Journey_Feedback__c'),
    ('Custom metadata', '3 types / 105 records', 'LoanSeva_Config__mdt (1), LoanSeva_FAQ__mdt (9), LoanSeva_Message__mdt (95, EN/HI/MR)'),
    ('Platform events', '2', 'LoanSeva_Event__e (registration), LoanSeva_Voice__e (voice notes)'),
    ('Apex classes', '104 (incl. tests)', 'Engine, actions, LOS services, sender, jobs, tests'),
    ('Triggers', '6', 'Event, Voice, ContentDocumentLink (inbound photos), Loan Application / Loan Journey / Existing Loan (INR stamp)'),
    ('Agentforce', '1 agent', 'Bot + BotVersion, genAiPlannerBundle, 2 genAiPlugins (topics), 39 genAiFunctions (actions)'),
    ('Lightning', '3 LWC, 1 app, 1 record page', 'loanSevaCreditReview, loanSevaScorecard, loanSevaDecisionModal; LoanSeva LOS app; Loan_Application_Record_Page'),
    ('Flows', '2', 'LoanSeva_WhatsApp_Inbound_Routing (Omni-Flow), LoanSeva_Send_Conversation_Message'),
    ('Security', '3 permission sets', 'LoanSeva_Guest (site), LoanSeva_Agent (agent user), LoanSeva_Admin (RMs/admins, also grants the ElevenLabs principal)'),
    ('Other', '', 'Queue LoanSeva_RM_Callbacks, Site LoanSeva, Named + External Credential LoanSeva_ElevenLabs, messaging components LoanSeva_Text and LoanSeva_Buttons_Test, static resource LoanSevaLogo (placeholder)'),
], widths=[3.2, 3.3, 10])

doc.add_heading('2.6 How a WhatsApp message flows (Bring Your Own Channel)', level=2)
para('WhatsApp reaches Salesforce through Bring Your Own Channel (BYOC). The channel definition in the org (SEs_WhatsAppBYOC) makes Salesforce the owner of consent and routing, '
     'and names the platform event that carries outbound messages: wpbyocdemo3__WP_BYOC_Meta_Event__e (payload in Payload__c, type in EventType__c). '
     'In production with Gupshup (Option 1a) the path is the same; Gupshup takes the place of the demo router.')
table([
    ('Step', 'From → to', 'What happens'),
    ('1', 'Customer → Meta', 'The customer sends text, a photo or a voice note on WhatsApp'),
    ('2', 'Meta → Gupshup', 'Meta calls the BSP\'s webhook'),
    ('3', 'Gupshup → Service Cloud', 'Gupshup calls the BYOC inbound API (Interaction Service API), authenticated with OAuth JWT via a connected app. Salesforce finds or creates the MessagingEndUser and the MessagingSession on the WhatsApp_BYOC channel'),
    ('4', 'Service Cloud → Agentforce', 'Omni-Channel runs the Omni-Flow LoanSeva_WhatsApp_Inbound_Routing, which routes the session to the LoanSeva by Kotak agent (fallback queue SDO_Service_Messaging)'),
    ('5', 'Agentforce → engine', 'The agent calls LoanSevaReplyAction with the customer\'s text'),
    ('6', 'Engine → Agentforce', 'LoanSevaEngine.reply() saves the journey, application, RM queue and LAN records and returns the reply and the next step; the agent relays it word for word'),
    ('7', 'Agentforce → Service Cloud', 'The reply is posted on the MessagingSession'),
    ('8', 'Service Cloud → Gupshup', 'Salesforce publishes the outbound message as the channel\'s platform event. Gupshup\'s connector subscribes via the Pub/Sub API (or a Salesforce flow calls Gupshup\'s send API). In the PoC the package flow wpbyocdemo3__WhatsApp_Send_Outbound_Event_to_Router_Prod calls the demo router'),
    ('9', 'Gupshup → Meta', 'Gupshup sends the message through its WhatsApp send API'),
    ('10', 'Meta → customer', 'The reply arrives in the same WhatsApp chat'),
], widths=[1.2, 4, 11.3])
bullets([
    ('One call per message.', 'Every customer message is a separate inbound API call (step 3) into the same open MessagingSession, so a 15-message chat makes 15 inbound calls and 15 or more outbound events. The session, the journey record and the transcript tie them together.'),
    ('Proactive messages', '(nudges, "your LAN is sanctioned") start at step 8: LoanSevaWhatsAppSender calls sendConversationMessages with the LoanSeva_Text messaging component.'),
    ('Human hand-off.', 'The agent escalates, Omni-Channel routes the session to the RM queue, and the RM accepts it in the Service Console. The RM\'s replies leave through steps 8–10, so the customer stays in the same chat.'),
    ('Photos and voice notes', 'arrive at step 3 as attachments, become Files on the session, and are picked up by the ContentDocumentLink trigger and the LoanSeva_Voice__e event.'),
    ('Hand back (Option 1a).', 'When the MessagingSession ends, Salesforce sends a conversation-closed event and Gupshup resumes its own bot. Gupshup has to build this routing and the BYOC connector; confirm whether they already have one.'),
    ('Option 1b instead.', 'Gupshup calls the Agentforce Agent API directly. There is no MessagingSession, no Omni-Channel and no outbound event: Gupshup sends the reply itself and keeps the human hand-off in its own inbox.'),
])

# ------------------------------------------------------------------ 3
doc.add_heading('3. How to enable it', level=1)
doc.add_heading('3.1 Prerequisites (licences and org features)', level=2)
bullets([
    ('Service Cloud with Messaging (Digital Engagement)', 'and Omni-Channel enabled.'),
    ('Agentforce / Einstein', 'enabled (Setup → Einstein Setup → Turn on Einstein; Agentforce Agents on), with an Einstein Agent user licence for the service agent.'),
    ('A WhatsApp channel:', 'either native WhatsApp in Messaging (Meta Cloud API) or Bring Your Own Channel for a BSP such as Gupshup. The demo uses the wpbyocdemo3 package.'),
    ('Sites', 'enabled with a registered Sites domain (for the registration page).'),
    ('Multi-currency', 'with INR active and set as corporate currency (Setup → Company Information → Currency Setup).'),
    ('Tooling:', 'Salesforce CLI (sf), Python 3 for the generator scripts.'),
])

doc.add_heading('3.2 Deploy the metadata', level=2)
code('sf project deploy start -d force-app -o <org> --dry-run        # validate\n'
     'scripts/deploy.sh force-app                                   # deploy\n'
     'sf apex run -f scripts/apex/create_agent_user.apex -o <org>    # once: Einstein Agent user\n'
     'sf apex run -f scripts/apex/post_deploy.apex -o <org>          # perm sets, RM queue member, drop-off job')
para('post_deploy.apex assigns LoanSeva_Guest to the site guest user, LoanSeva_Agent to the agent user and LoanSeva_Admin to the running admin, adds the admin to the LoanSeva_RM_Callbacks queue, and schedules LoanSevaDropOffJob.')
note('In another org, update the serviceChannelId and botId inside LoanSeva_WhatsApp_Inbound_Routing.flow-meta.xml before deploying; they are kotak-demo record ids.')

doc.add_heading('3.3 Agent', level=2)
para('The agent is deployed as classic Agentforce metadata and activated from the CLI. To change it: regenerate, deactivate, deploy, activate.')
code('PYTHONPATH=scripts python3 scripts/gen_classic_agent.py\n'
     'sf agent deactivate --api-name LoanSeva_by_Kotak -o <org>\n'
     'sf project deploy start -d force-app/main/default/genAiFunctions -d force-app/main/default/genAiPlugins \\\n'
     '   -d force-app/main/default/genAiPlannerBundles -d force-app/main/default/bots -o <org>\n'
     'sf agent activate --api-name LoanSeva_by_Kotak -o <org>')
note('Redeploying the agent breaks open WhatsApp conversations. End every Active Messaging Session (Service Console → End Conversation) before deactivating, and have the customer send a fresh "Hi" afterwards.')

doc.add_heading('3.4 WhatsApp channel and routing', level=2)
steps([
    ('Point the channel at the flow:', 'run scripts/route_whatsapp.sh (sets Omni-Flow routing to LoanSeva_WhatsApp_Inbound_Routing, fallback queue SDO_Service_Messaging). scripts/route_whatsapp.sh --queue rolls back.'),
    ('Activate the messaging component:', 'Setup → Messaging Components → LoanSeva Text → Activate. Deployed components arrive inactive; without this, sends fail with MessagingComponentNotActive.'),
    ('Register presenter numbers', 'with the WhatsApp router (demo package only): WhatsApp Demo Component tab → add each number with country code against this org. Then send "Hi" from each phone to open the 24-hour window.'),
    ('WhatsApp template:', 'for business-initiated first messages, create and get approval for a Utility template in Meta Business Manager (or in Gupshup) and map it in the channel settings.'),
    ('Live hand-off:', 'the human must be in the fallback queue. In the Service Console open Omni-Channel, set status to Messaging - Available and accept incoming chats.'),
])

doc.add_heading('3.5 Registration site', level=2)
steps([
    'Setup → Sites → LoanSeva: confirm Active, home page LoanSevaRegister.',
    'Public Access Settings: Visualforce page LoanSevaRegister and Apex class LoanSevaRegisterController enabled (the deploy does this).',
    'Test https://<sites-domain>/loanseva?src=qr. Regenerate the QR if the domain changes (command in SETUP_STEPS.md §2).',
])
note('In the kotak-demo trial org the public site returns 503 at the org level. Fallback: logged-in users open /apex/LoanSevaRegister?src=link.')

doc.add_heading('3.6 Configuration', level=2)
table([
    ('Setting (LoanSeva_Config__mdt.Default)', 'Demo value', 'Purpose'),
    ('Send_Mode__c', 'ConversationMessage', 'Real WhatsApp sends; LogOnly composes and logs without sending'),
    ('RM_Owner_Username__c', 'demo admin', 'Owner of RM Tasks and Events'),
    ('Demo_OTP__c', '123456', 'Fixed OTP; replace with a real OTP provider in production'),
    ('Idle / Nudge 1 / Nudge 2 minutes', '2 / 2 / 5', 'Drop-off timings (production 1440 / 4320)'),
    ('Voice_Enabled__c', 'true', 'Turns voice-note transcription on or off'),
    ('Rates, FOIR, fees, tenures', 'Illustrative', 'Align with Kotak published terms before external use'),
], widths=[5.5, 3.5, 7.5])
para('Edit values in scripts/content.py, regenerate and redeploy the custom metadata.')

doc.add_heading('3.7 Voice notes (ElevenLabs)', level=2)
para('Setup → Named Credentials → External Credentials → LoanSeva_ElevenLabs → Principal Default → set the authentication parameter ApiKey to an ElevenLabs key with Speech to Text permission. The voice subscriber user needs LoanSeva_Admin.')

doc.add_heading('3.8 LOS access for RMs', level=2)
bullets([
    'Assign LoanSeva_Admin to each RM; the LoanSeva LOS app and Loan Application record page are assigned in metadata (no App Builder activation).',
    'Add RMs to the LoanSeva_RM_Callbacks queue (Setup → Queues) so they see the RM Callback Queue list view.',
    'Users may need a hard refresh (Cmd+Shift+R) after a deploy, because the browser caches the LWCs.',
])

doc.add_heading('3.9 Demo data and verification', level=2)
code('sf apex run -f scripts/apex/seed.apex -o <org>     # demo customers; edit DEMO_MOBILES first; re-run to reset\n'
     'python3 scripts/runapex.py scripts/apex/smoke.apex # full journey without WhatsApp, prints every reply\n'
     'sf apex run test -o <org> -n LoanSevaEngineTest -n LoanSevaActionsTest ... -w 20   # all 14 test classes')
table([
    ('Demo customer', 'Existing loan', 'Story'),
    ('Priya Sharma', 'EMI ₹11,800, 12/12 on time, pre-approved ₹5L', 'Happy-path Top-Up at the best rate'),
    ('Rahul Deshmukh', 'EMI ₹8,600, 11/12 on time, pre-approved ₹3L', 'Good track: rate +1%, 80% cap; also new-loan demo with a PAN'),
    ('Ananya Iyer', 'EMI ₹21,500, 12/12 on time, pre-approved ₹10L', 'High value; low income shows FOIR capping the offer'),
], widths=[3.5, 7, 6])

# ------------------------------------------------------------------ 4
doc.add_heading('4. Known limitations and path to production', level=1)
table([
    ('Area', 'In the PoC', 'For production'),
    ('WhatsApp channel', 'Demo BYOC package and shared router', 'Recommended: a dedicated loans number on Salesforce native WhatsApp (alternative: Gupshup via BYOC); approved templates'),
    ('OTP', 'Fixed demo OTP', 'OTP provider behind a Named Credential'),
    ('Bank data', 'Seeded Existing_Loan__c records', 'Core banking / loan servicing via MuleSoft'),
    ('Underwriting', 'Mock NSDL, CIBIL, Perfios, EPFO, Hunter', 'Real provider integrations and Kotak credit policy'),
    ('Attachments in agent chat', 'Photos arrive as Files on the session and are linked by trigger', 'AiAgentFileUploadsPilot / AgentAttachmentsPilot permissions must be enabled by Salesforce for the org'),
    ('Public site', 'Trial org site returns 503', 'Production org Sites / Experience Cloud domain'),
    ('Agent changes', 'Redeploy ends open conversations', 'Release in a maintenance window; version the agent'),
    ('Personalisation', 'Fixed nudge timers', 'Data Cloud segments and next best action'),
    ('Branding & content', 'Placeholder logo; illustrative rates and FAQ', 'Official assets; compliance-approved copy'),
], widths=[3.3, 5.6, 7.6])

for s in doc.sections:
    f = s.footer.paragraphs[0]
    f.text = 'LoanSeva by Kotak · Build and setup guide'
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f.runs[0].font.size = Pt(8)
    f.runs[0].font.color.rgb = GREY

out = 'docs/LoanSeva_Build_and_Setup_Guide.docx'
doc.save(out)
print(out)
