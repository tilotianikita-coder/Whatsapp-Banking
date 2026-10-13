"""Builds docs/LoanSeva_WhatsApp_Cost_Model.xlsx: options 1a / 1b / 2 side by side. Run: python3 scripts/gen_cost_model.py

Prices are left blank on purpose (fill in from the Salesforce quote and Gupshup's contract); volumes are example values.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY, RED, SF = '003874', 'ED1C24', '0176D3'
INPUT = PatternFill('solid', fgColor='FFF4CC')   # yellow: volumes
PRICE = PatternFill('solid', fgColor='E5F1FB')   # blue: prices to fill in
HEAD = PatternFill('solid', fgColor=NAVY)
SUB = PatternFill('solid', fgColor='F3F4F6')
TOTAL = PatternFill('solid', fgColor='E8F5EC')
thin = Side(style='thin', color='C9CDD2')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
INR = '"₹"#,##0'
INR2 = '"₹"#,##0.00'

wb = Workbook()

# ------------------------------------------------------------------ Inputs
ws = wb.active
ws.title = 'Inputs'
ws.column_dimensions['A'].width = 52
ws.column_dimensions['B'].width = 16
ws.column_dimensions['C'].width = 14
ws.column_dimensions['D'].width = 62

ws['A1'] = 'LoanSeva WhatsApp banking: cost model inputs'
ws['A1'].font = Font(bold=True, size=16, color=NAVY)
ws['A2'] = 'Yellow = volumes (example values, replace with Kotak\'s). Blue = prices (blank: fill in from the Salesforce quote and the Gupshup contract). All amounts in INR.'
ws['A2'].font = Font(italic=True, color='5F6368')
ws.merge_cells('A2:D2')

names = {}
row = 4


def section(title):
    global row
    for c, v in enumerate((title, 'Value', 'Unit', 'Notes'), start=1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEAD
    row += 1


def item(key, label, value, unit, note, fill, fmt=None):
    global row
    ws.cell(row=row, column=1, value=label).border = BOX
    v = ws.cell(row=row, column=2, value=value)
    v.fill = fill
    v.border = BOX
    v.number_format = fmt or ('#,##0' if fill is INPUT else INR)
    ws.cell(row=row, column=3, value=unit).border = BOX
    n = ws.cell(row=row, column=4, value=note)
    n.border = BOX
    n.alignment = Alignment(wrap_text=True, vertical='top')
    names[key] = f"Inputs!$B${row}"
    row += 1


section('Volumes (loan journey only)')
item('months', 'Contract term', 12, 'months', 'Used to annualise monthly costs', INPUT)
item('rm_los', 'RMs using the LOS and RM callback queue', 50, 'users', 'Need a Service Cloud / Financial Services Cloud seat in every option', INPUT)
item('rm_chat', 'Of those, RMs who take live WhatsApp chats', 20, 'users', 'Need the messaging add-on (1a, 2) or a Gupshup inbox seat (1b)', INPUT)
item('chats', 'WhatsApp loan conversations per month', 20000, 'chats', 'Every conversation the Agentforce agent handles', INPUT)
item('actions', 'Agent actions per conversation', 12, 'actions', 'LoanSeva calls one action per customer reply; a Top-Up is roughly 10-15 replies', INPUT)
item('mkt', 'Paid marketing templates per month', 30000, 'messages', 'Top-Up campaign invites (business-initiated)', INPUT)
item('util', 'Paid utility templates per month', 8000, 'messages', 'Nudges / updates sent outside the 24-hour customer window', INPUT)
item('auth', 'Paid authentication templates per month', 0, 'messages', 'OTP over WhatsApp, if used (PoC uses a demo OTP)', INPUT)
row += 1

section('Salesforce prices')
item('p_seat', 'Service Cloud / FSC licence', None, 'per user / month', 'Edition per the quote (e.g. FSC Enterprise / Unlimited)', PRICE)
item('p_msg', 'Messaging add-on (Digital Engagement)', None, 'per user / month', 'Options 1a and 2 only; confirm SKU and whether BYOC is included', PRICE)
item('p_byoc', 'BYOC usage charge, if any', None, 'per conversation', 'Option 1a only; confirm whether BYOC is metered', PRICE, INR2)
item('p_action', 'Agentforce price per action', None, 'per action', 'Flex Credits rate per the quote; use per-conversation pricing by setting actions = 1', PRICE, INR2)
item('p_sf_msg', 'Salesforce messaging charge (excl. Meta fee)', None, 'per paid message', 'Option 2 only; Salesforce message pack price minus the Meta pass-through', PRICE, INR2)
item('p_mule', 'MuleSoft', None, 'per year', 'Core banking, bureau, KYC, Perfios, e-sign APIs', PRICE)
item('p_dc', 'Data Cloud', None, 'per year', 'Phase 3; set 0 to exclude', PRICE)
item('p_shield', 'Shield (encryption, event monitoring)', None, 'per year', 'Usually required by banks', PRICE)
item('p_exp', 'Experience Cloud (registration page)', None, 'per year', 'Small; set 0 if covered elsewhere', PRICE)
row += 1

section('Meta WhatsApp fees (same in every option)')
item('m_mkt', 'Meta fee: marketing template', None, 'per message', 'Meta India rate card', PRICE, INR2)
item('m_util', 'Meta fee: utility template', None, 'per message', 'Free inside the 24-hour customer window', PRICE, INR2)
item('m_auth', 'Meta fee: authentication template', None, 'per message', 'Meta India rate card', PRICE, INR2)
row += 1

section('Gupshup prices (loan traffic only)')
item('g_markup', 'Gupshup markup on top of the Meta fee', None, 'per paid message', 'Options 1a and 1b; option 2 loan traffic bypasses Gupshup', PRICE, INR2)
item('g_platform', 'Gupshup platform fee for the loan journey', None, 'per month', 'Incremental fee for loans traffic; the main number\'s existing fees are the same in all options and left out', PRICE)
item('g_handover', 'Gupshup hand-over to Salesforce', None, 'per month', 'Option 1a: forwarding loan conversations via BYOC', PRICE)
item('g_seat', 'Gupshup agent-inbox seat', None, 'per user / month', 'Option 1b: RMs take live chats in Gupshup, not Salesforce', PRICE)
row += 1

section('One-off costs')
item('o_1a', 'Option 1a: Gupshup hand-over / hand-back build', None, 'one-off', 'Loan-intent routing to BYOC and back', PRICE)
item('o_1b', 'Option 1b: Gupshup Agent API integration build', None, 'one-off', 'Gupshup calls the Agentforce Agent API', PRICE)
item('o_2', 'Option 2: second number setup', None, 'one-off', 'Number, display-name verification, template approvals, launch comms', PRICE)
item('o_impl', 'Salesforce implementation (all options)', None, 'one-off', 'Agent, journeys, LOS, integrations; same in every option', PRICE)

# ------------------------------------------------------------------ Comparison
cs = wb.create_sheet('Comparison')
for col, w in zip('ABCDEF', (50, 18, 18, 18, 12, 58)):
    cs.column_dimensions[col].width = w
cs['A1'] = 'Annual cost by option (loan journey)'
cs['A1'].font = Font(bold=True, size=16, color=NAVY)
cs['A2'] = 'Calculated from the Inputs sheet. Shows 0 until prices are filled in.'
cs['A2'].font = Font(italic=True, color='5F6368')
hdr = ('Cost line', '1a: Gupshup + BYOC', '1b: Gupshup + Agent API', '2: Salesforce native', 'Billed by', 'How it is calculated')
for c, v in enumerate(hdr, start=1):
    cell = cs.cell(row=4, column=c, value=v)
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = HEAD
    cell.alignment = Alignment(wrap_text=True, vertical='center')
cs.row_dimensions[4].height = 32

N = names
paid = f"({N['mkt']}*{N['m_mkt']}+{N['util']}*{N['m_util']}+{N['auth']}*{N['m_auth']})"
count = f"({N['mkt']}+{N['util']}+{N['auth']})"
lines = [
    ('Salesforce', None),
    ('RM seats (Service Cloud / FSC)', [f"={N['rm_los']}*{N['p_seat']}*{N['months']}"] * 3, 'Salesforce', 'RMs using the LOS × licence × months'),
    ('Messaging add-on for live-chat RMs', [f"={N['rm_chat']}*{N['p_msg']}*{N['months']}", '=0', f"={N['rm_chat']}*{N['p_msg']}*{N['months']}"], 'Salesforce', 'Live-chat RMs × add-on × months (not needed in 1b)'),
    ('BYOC usage', [f"={N['chats']}*{N['p_byoc']}*{N['months']}", '=0', '=0'], 'Salesforce', 'Conversations × BYOC charge × months (1a only, if metered)'),
    ('Agentforce usage', [f"={N['chats']}*{N['actions']}*{N['p_action']}*{N['months']}"] * 3, 'Salesforce', 'Conversations × actions × price per action × months'),
    ('Salesforce messaging charge (loan number)', ['=0', '=0', f"={count}*{N['p_sf_msg']}*{N['months']}"], 'Salesforce', 'Paid messages × Salesforce charge × months (option 2 only)'),
    ('MuleSoft', [f"={N['p_mule']}*{N['months']}/12"] * 3, 'Salesforce', 'Annual price, pro rata to the term'),
    ('Data Cloud', [f"={N['p_dc']}*{N['months']}/12"] * 3, 'Salesforce', 'Annual price, pro rata to the term'),
    ('Shield', [f"={N['p_shield']}*{N['months']}/12"] * 3, 'Salesforce', 'Annual price, pro rata to the term'),
    ('Experience Cloud', [f"={N['p_exp']}*{N['months']}/12"] * 3, 'Salesforce', 'Annual price, pro rata to the term'),
    ('Gupshup', None),
    ('Gupshup markup on paid messages', [f"={count}*{N['g_markup']}*{N['months']}"] * 2 + ['=0'], 'Gupshup', 'Paid messages × markup × months (loan traffic leaves Gupshup in option 2)'),
    ('Gupshup platform fee (loan journey)', [f"={N['g_platform']}*{N['months']}"] * 2 + ['=0'], 'Gupshup', 'Monthly fee × months'),
    ('Gupshup hand-over to Salesforce', [f"={N['g_handover']}*{N['months']}", '=0', '=0'], 'Gupshup', 'Monthly fee × months (1a only)'),
    ('Gupshup agent-inbox seats', ['=0', f"={N['rm_chat']}*{N['g_seat']}*{N['months']}", '=0'], 'Gupshup', 'Live-chat RMs × seat × months (1b only)'),
    ('Meta', None),
    ('Meta WhatsApp fees', [f"={paid}*{N['months']}"] * 3, 'Meta via provider', 'Same in every option; passed through by Gupshup (1a, 1b) or Salesforce (2)'),
]
r = 5
sf_rows, gs_rows, meta_rows = [], [], []
bucket = None
for line in lines:
    if line[1] is None:
        bucket = line[0]
        cell = cs.cell(row=r, column=1, value=line[0])
        cell.font = Font(bold=True, color=NAVY)
        for c in range(1, 7):
            cs.cell(row=r, column=c).fill = SUB
        r += 1
        continue
    label, formulas, by, how = line
    cs.cell(row=r, column=1, value=label).border = BOX
    for c, f in enumerate(formulas, start=2):
        cell = cs.cell(row=r, column=c, value=f)
        cell.number_format = INR
        cell.border = BOX
    cs.cell(row=r, column=5, value=by).border = BOX
    h = cs.cell(row=r, column=6, value=how)
    h.border = BOX
    h.font = Font(color='5F6368', size=9)
    {'Salesforce': sf_rows, 'Gupshup': gs_rows, 'Meta': meta_rows}[bucket].append(r)
    r += 1

r += 1


def total(label, rows_, bold=False, fill=TOTAL):
    global r
    cs.cell(row=r, column=1, value=label).font = Font(bold=True, color=NAVY if bold else '1F2328')
    for c, col in enumerate('BCD', start=2):
        f = '=' + '+'.join(f'{col}{x}' for x in rows_) if rows_ else '=0'
        cell = cs.cell(row=r, column=c, value=f)
        cell.number_format = INR
        cell.font = Font(bold=bold)
        cell.fill = fill
        cell.border = BOX
    r += 1
    return r - 1


t_sf = total('Salesforce recurring', sf_rows)
t_gs = total('Gupshup recurring', gs_rows)
t_meta = total('Meta fees', meta_rows)
t_all = r
cs.cell(row=r, column=1, value='Total recurring for the term').font = Font(bold=True, color=NAVY, size=12)
for c, col in enumerate('BCD', start=2):
    cell = cs.cell(row=r, column=c, value=f'={col}{t_sf}+{col}{t_gs}+{col}{t_meta}')
    cell.number_format = INR
    cell.font = Font(bold=True, size=12)
    cell.fill = TOTAL
    cell.border = BOX
r += 1
t_one = r
cs.cell(row=r, column=1, value='One-off costs').font = Font(bold=True)
for c, opt in zip((2, 3, 4), ('o_1a', 'o_1b', 'o_2')):
    cell = cs.cell(row=r, column=c, value=f"={N[opt]}+{N['o_impl']}")
    cell.number_format = INR
    cell.border = BOX
r += 1
cs.cell(row=r, column=1, value='Year-1 total (recurring + one-off)').font = Font(bold=True, color=RED, size=12)
for c, col in enumerate('BCD', start=2):
    cell = cs.cell(row=r, column=c, value=f'={col}{t_all}+{col}{t_one}')
    cell.number_format = INR
    cell.font = Font(bold=True, size=12, color=RED)
    cell.border = BOX
r += 2

cs.cell(row=r, column=1, value='Share of the recurring spend').font = Font(bold=True, color=NAVY)
r += 1
for label, t in (('Salesforce share', t_sf), ('Gupshup share', t_gs), ('Meta share', t_meta)):
    cs.cell(row=r, column=1, value=label)
    for c, col in enumerate('BCD', start=2):
        cell = cs.cell(row=r, column=c, value=f'=IF({col}{t_all}=0,0,{col}{t}/{col}{t_all})')
        cell.number_format = '0%'
        cell.border = BOX
    r += 1
cs.cell(row=r, column=1, value='Cost per loan conversation (recurring)')
for c, col in enumerate('BCD', start=2):
    cell = cs.cell(row=r, column=c, value=f"=IF({N['chats']}*{N['months']}=0,0,{col}{t_all}/({N['chats']}*{N['months']}))")
    cell.number_format = INR2
    cell.border = BOX
r += 2

cs.cell(row=r, column=1, value='What you get for the money').font = Font(bold=True, color=NAVY)
r += 1
for label, vals in (
    ('RM live takeover in the Salesforce console, with history', ('Yes', 'No (Gupshup inbox)', 'Yes')),
    ('Conversation transcript on the customer record', ('Yes', 'No', 'Yes')),
    ('Same WhatsApp number as today', ('Yes', 'Yes', 'No (second number)')),
    ('Vendors in the loan conversation', ('2', '2', '1')),
):
    cs.cell(row=r, column=1, value=label).border = BOX
    for c, v in enumerate(vals, start=2):
        cell = cs.cell(row=r, column=c, value=v)
        cell.border = BOX
        cell.alignment = Alignment(horizontal='center')
    r += 1
cs.freeze_panes = 'B5'

# ------------------------------------------------------------------ Notes
ns = wb.create_sheet('Notes')
ns.column_dimensions['A'].width = 120
notes = [
    ('How to use', True),
    ('1. Replace the yellow volumes on Inputs with Kotak\'s numbers.', False),
    ('2. Fill the blue prices from the Salesforce quote (AE / pricing desk), the Gupshup contract and Meta\'s India rate card.', False),
    ('3. Read the totals, the Salesforce / Gupshup / Meta split and the cost per conversation on Comparison.', False),
    ('', False),
    ('The three options', True),
    ('1a  Gupshup keeps the number; loan conversations are forwarded to Salesforce via Bring Your Own Channel. RMs take over in the Salesforce console.', False),
    ('1b  Gupshup keeps the number and calls the Agentforce Agent API directly. No Salesforce messaging add-on; human hand-off stays in Gupshup\'s inbox.', False),
    ('2   A second "Kotak Loans" number on Salesforce native WhatsApp. Loan traffic bypasses Gupshup; the main number stays on Gupshup.', False),
    ('', False),
    ('Assumptions', True),
    ('Only the loan journey is modelled. The main number\'s other traffic (alerts, OTPs, campaigns) and its Gupshup fees are the same in every option, so they are left out.', False),
    ('Meta charges per message. Replies inside the 24-hour window after the customer writes are largely free, so only templates (marketing, utility outside the window, authentication) are counted as paid.', False),
    ('Agentforce is modelled per action. For per-conversation pricing, set "Agent actions per conversation" to 1 and enter the per-conversation price.', False),
    ('Licence names, metering and Meta rates change; confirm the current SKUs and prices before sharing figures with Kotak.', False),
]
for i, (t, bold) in enumerate(notes, start=1):
    c = ns.cell(row=i, column=1, value=t)
    c.font = Font(bold=bold, color=NAVY if bold else '1F2328', size=12 if bold else 11)
    c.alignment = Alignment(wrap_text=True)

out = 'docs/LoanSeva_WhatsApp_Cost_Model.xlsx'
wb.save(out)
print(out)
