"""Builds docs/LoanSeva_WhatsApp_Banking_Why_Salesforce.pptx (10 slides, 16:9). Run: python3 scripts/gen_deck.py"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

RED = RGBColor(0xED, 0x1C, 0x24)
NAVY = RGBColor(0x00, 0x38, 0x74)
SF = RGBColor(0x01, 0x76, 0xD3)
SF_LIGHT = RGBColor(0xE5, 0xF1, 0xFB)
GREY = RGBColor(0x5F, 0x63, 0x68)
LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
LINE = RGBColor(0xC9, 0xCD, 0xD2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1F, 0x23, 0x28)
GREEN = RGBColor(0x2E, 0x84, 0x4A)
GREEN_LIGHT = RGBColor(0xE8, 0xF5, 0xEC)
AMBER_LIGHT = RGBColor(0xFE, 0xF3, 0xE2)
RED_LIGHT = RGBColor(0xFD, 0xEC, 0xEC)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, runs, size=14, color=DARK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """runs: a string, or a list of paragraphs; each paragraph a string or (text, {size,color,bold,bullet})."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        t, o = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = o.get('align', align)
        para.space_after = Pt(o.get('after', 4))
        r = para.add_run()
        r.text = ('•  ' if o.get('bullet') else '') + t
        r.font.size = Pt(o.get('size', size))
        r.font.bold = o.get('bold', bold)
        r.font.color.rgb = o.get('color', color)
        r.font.name = 'Calibri'
    return tb


def box(slide, x, y, w, h, label, fill=WHITE, line=LINE, color=DARK, size=12, bold=False, sub=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, sub_size=10):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.shadow.inherit = False
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.12
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ('margin_left', 'margin_right'):
        setattr(tf, m, Inches(0.08))
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = 'Calibri'
    if sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = sub
        r2.font.size = Pt(sub_size)
        r2.font.color.rgb = GREY if fill == WHITE or fill == LIGHT else color
        r2.font.name = 'Calibri'
    return s


def arrow(slide, x1, y1, x2, y2, color=GREY, dashed=False, both=False, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dashed:
        etree.SubElement(ln, qn('a:prstDash'), val='dash')
    if both:
        etree.SubElement(ln, qn('a:headEnd'), type='triangle', w='med', len='med')
    etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='med', len='med')
    return c


def header(slide, title, kicker, n):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RED
    bar.line.fill.background()
    text(slide, Inches(0.6), Inches(0.35), Inches(12), Inches(0.35), kicker.upper(), size=12, color=RED, bold=True)
    text(slide, Inches(0.6), Inches(0.65), Inches(12.2), Inches(0.8), title, size=28, color=NAVY, bold=True)
    text(slide, Inches(0.6), Inches(7.05), Inches(9), Inches(0.3), 'LoanSeva by Kotak · WhatsApp Banking on Salesforce', size=10, color=GREY)
    text(slide, Inches(12.0), Inches(7.05), Inches(0.8), Inches(0.3), str(n), size=10, color=GREY, align=PP_ALIGN.RIGHT)


def group(slide, x, y, w, h, title, fill, line, title_color):
    g = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    g.adjustments[0] = 0.04
    g.shadow.inherit = False
    g.fill.solid()
    g.fill.fore_color.rgb = fill
    g.line.color.rgb = line
    g.line.width = Pt(1.25)
    text(slide, x + Inches(0.15), y + Inches(0.06), w - Inches(0.3), Inches(0.3), title, size=11, color=title_color, bold=True)
    return g


# ---------------------------------------------------------------- 1. Title
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = NAVY
bg.line.fill.background()
stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(5.1), prs.slide_width, Inches(0.08))
stripe.fill.solid()
stripe.fill.fore_color.rgb = RED
stripe.line.fill.background()
text(s, Inches(0.8), Inches(1.2), Inches(11.5), Inches(0.4), 'KOTAK MAHINDRA BANK  ×  SALESFORCE', size=14, color=RGBColor(0xFF, 0xB3, 0xB6), bold=True)
text(s, Inches(0.8), Inches(1.7), Inches(11.5), Inches(1.6), 'From a WhatsApp bot to a banker in every chat', size=44, color=WHITE, bold=True)
text(s, Inches(0.8), Inches(3.25), Inches(11.5), Inches(1.6), [
    ('Why Salesforce for conversational AI on WhatsApp, and why running it natively is the simplest path', {'size': 20, 'color': RGBColor(0xD6, 0xE2, 0xF0)}),
], size=20)
text(s, Inches(0.8), Inches(5.45), Inches(11.8), Inches(1.4), [
    ('The ask in one line', {'size': 13, 'color': RGBColor(0xFF, 0xB3, 0xB6), 'bold': True}),
    ('Run the loan journey natively on Salesforce: WhatsApp channel, AI agent and LOS on one platform, where the customer, the loan and the RM already are. Least to build, one vendor.', {'size': 16, 'color': WHITE}),
])

# ---------------------------------------------------------------- 2. Challenge
s = prs.slides.add_slide(BLANK)
header(s, 'Today the WhatsApp bot talks, but it doesn\'t know the customer', 'The challenge today', 2)
cards = [
    ('No customer context', 'The bot sees a phone number, not a relationship: existing loans, EMI track, pre-approved offers, open applications and past conversations sit in other systems.'),
    ('Scripted menus, not conversations', 'Decision trees break on free text, Hinglish, voice notes or a question asked mid-journey ("what\'s the EMI on 3 lakh?"). Customers drop out or ask for an agent.'),
    ('Hand-off loses the thread', 'When a human takes over, they start from zero: no journey state, no offer, no documents. The customer repeats themselves.'),
    ('Journeys end in a lead, not a loan', 'Answers aren\'t written into the LOS. Underwriting, RM callback and sanction happen in separate tools, re-keyed by hand.'),
    ('Drop-offs go cold', 'A customer who stops at step 4 isn\'t followed up. There\'s no RM task, no nudge and no record of where or why they stopped.'),
    ('Fragmented view and governance', 'Bot analytics, CRM and LOS reports don\'t join up. Prompts, consent, PII masking and audit are handled differently in each tool.'),
]
cw, ch = Inches(3.95), Inches(2.3)
for i, (t, d) in enumerate(cards):
    x = Inches(0.6) + (i % 3) * (cw + Inches(0.2))
    y = Inches(1.65) + (i // 3) * (ch + Inches(0.2))
    c = box(s, x, y, cw, ch, '', fill=LIGHT, line=None)
    acc = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + Inches(0.2), Inches(0.07), Inches(0.5))
    acc.fill.solid()
    acc.fill.fore_color.rgb = RED
    acc.line.fill.background()
    text(s, x + Inches(0.25), y + Inches(0.18), cw - Inches(0.4), Inches(0.5), t, size=17, color=NAVY, bold=True)
    text(s, x + Inches(0.25), y + Inches(0.72), cw - Inches(0.4), ch - Inches(0.8), d, size=13, color=DARK)
text(s, Inches(0.6), Inches(6.62), Inches(12.2), Inches(0.4),
     'The gap is not the WhatsApp channel. Gupshup delivers messages well. The gap is context, intelligence and the link into the loan process.',
     size=13, color=RED, bold=True)

# ---------------------------------------------------------------- 3. Current state
s = prs.slides.add_slide(BLANK)
header(s, 'Current state: the channel and the bank are separate islands', 'Current-state architecture (as understood; to validate with Kotak)', 3)
# Customer
box(s, Inches(0.5), Inches(3.2), Inches(1.6), Inches(1.1), 'Customer', fill=GREEN_LIGHT, line=GREEN, bold=True, sub='WhatsApp')
# Meta
box(s, Inches(2.55), Inches(3.2), Inches(1.5), Inches(1.1), 'Meta', fill=WHITE, bold=True, sub='WhatsApp Business Platform')
# Gupshup group
group(s, Inches(4.5), Inches(1.65), Inches(3.4), Inches(4.3), 'GUPSHUP (BSP + BOT)', AMBER_LIGHT, RGBColor(0xE0, 0x9B, 0x2D), RGBColor(0x9A, 0x5B, 0x00))
box(s, Inches(4.75), Inches(2.1), Inches(2.9), Inches(0.8), 'Message API & templates', sub='Delivery, opt-in, campaigns')
box(s, Inches(4.75), Inches(3.05), Inches(2.9), Inches(0.8), 'Bot builder', sub='Menus, keyword flows, FAQ')
box(s, Inches(4.75), Inches(4.0), Inches(2.9), Inches(0.8), 'Agent inbox', sub='Separate console for live chat')
box(s, Inches(4.75), Inches(4.95), Inches(2.9), Inches(0.8), 'Bot analytics', sub='Sessions, clicks, drop-offs')
# Bank systems
group(s, Inches(8.5), Inches(1.65), Inches(4.3), Inches(4.3), 'BANK SYSTEMS (EACH ITS OWN SILO)', LIGHT, LINE, GREY)
items = [('Core banking / loan servicing', 'Existing loans, EMIs, track'), ('LOS', 'Applications, underwriting'),
         ('CRM / RM tools', 'Leads, tasks, callbacks'), ('Bureau · KYC · Perfios', 'Called from the LOS')]
for i, (t, d) in enumerate(items):
    box(s, Inches(8.75), Inches(2.1) + i * Inches(0.95), Inches(3.8), Inches(0.8), t, sub=d)
arrow(s, Inches(2.1), Inches(3.75), Inches(2.55), Inches(3.75), both=True)
arrow(s, Inches(4.05), Inches(3.75), Inches(4.75), Inches(3.45), both=True)
# point-to-point integrations, dashed
arrow(s, Inches(7.65), Inches(3.45), Inches(8.75), Inches(2.5), color=RED, dashed=True)
arrow(s, Inches(7.65), Inches(4.4), Inches(8.75), Inches(4.4), color=RED, dashed=True)
text(s, Inches(7.55), Inches(5.95), Inches(5.3), Inches(0.5), 'Dashed red: point-to-point API calls or CSV exports, built per use case', size=10, color=RED)
# pain callouts
pains = ['Bot can\'t see the loan, offer or journey', 'Live agent starts with no history', 'Answers re-keyed into the LOS', 'No single view of a drop-off']
for i, p in enumerate(pains):
    x = Inches(0.5) + i * Inches(3.1)
    box(s, x, Inches(6.4), Inches(2.95), Inches(0.5), p, fill=RED_LIGHT, line=None, color=RGBColor(0xB0, 0x10, 0x18), size=11, bold=True)

# ---------------------------------------------------------------- 4. Why Salesforce
s = prs.slides.add_slide(BLANK)
header(s, 'Why Salesforce for conversational AI: context is the product', 'Why Salesforce', 4)
rows = [
    ('Capability', 'Bot on a messaging platform', 'Agentforce on Salesforce'),
    ('Customer context', 'Phone number plus what the bot is passed through APIs', 'Native: contact, existing loans, offers, applications, cases and past chats on one record'),
    ('Conversation', 'Menus and keyword flows; LLM add-ons sit outside bank data', 'LLM agent grounded in bank data; a deterministic Apex engine keeps the numbers and rules exact'),
    ('Human hand-off', 'Separate agent inbox', 'Omni-Channel routes to the right RM queue with the full transcript and journey state'),
    ('Loan process', 'Ends in a lead or a webhook', 'The chat creates the application, runs checks, books the RM callback and sanctions the LAN'),
    ('Languages & voice', 'Per-flow translations', 'English, Hindi, Marathi and voice notes on the same journey; switch any time'),
    ('Trust & governance', 'Bot vendor plus each integration', 'Einstein Trust Layer (masking, zero retention), sharing, field security, audit, consent'),
    ('Insight & re-engagement', 'Bot analytics only', 'Drop-off point on the record; RM task and nudges; Data Cloud segments for next best action'),
]
tw = [Inches(2.3), Inches(4.4), Inches(5.4)]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.6), Inches(1.6), sum(tw, Emu(0)), Inches(4.9)).table
for c, w in enumerate(tw):
    tbl.columns[c].width = w
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.margin_left = cell.margin_right = Inches(0.1)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.fill.solid()
        if r == 0:
            cell.fill.fore_color.rgb = NAVY if c < 2 else SF
        else:
            cell.fill.fore_color.rgb = (LIGHT if c < 2 else SF_LIGHT) if r % 2 else WHITE
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = val
        run.font.name = 'Calibri'
        run.font.size = Pt(13 if r == 0 else 12)
        run.font.bold = r == 0 or c == 0
        run.font.color.rgb = WHITE if r == 0 else (NAVY if c == 0 else DARK)
text(s, Inches(0.6), Inches(6.6), Inches(12.2), Inches(0.4),
     'Not either/or: Salesforce can run behind Gupshup on the same number, or be the WhatsApp provider itself (native, direct with Meta).',
     size=13, color=RED, bold=True)

# ---------------------------------------------------------------- 5. Future state
s = prs.slides.add_slide(BLANK)
header(s, 'Future state: one platform from "Hi" to sanctioned loan', 'Future-state architecture with Salesforce', 5)
box(s, Inches(0.4), Inches(2.0), Inches(1.45), Inches(1.0), 'Customer', fill=GREEN_LIGHT, line=GREEN, bold=True, sub='WhatsApp: text, voice, photos')
box(s, Inches(0.4), Inches(4.3), Inches(1.45), Inches(1.0), 'QR / SMS / email', fill=GREEN_LIGHT, line=GREEN, bold=True, sub='Campaign entry')
group(s, Inches(2.15), Inches(1.65), Inches(1.75), Inches(4.0), 'CHANNEL', AMBER_LIGHT, RGBColor(0xE0, 0x9B, 0x2D), RGBColor(0x9A, 0x5B, 0x00))
box(s, Inches(2.3), Inches(2.05), Inches(1.45), Inches(1.6), 'WhatsApp', bold=True, sub='A: via Gupshup (BYOC) · B and end state: Salesforce native')
box(s, Inches(2.3), Inches(3.85), Inches(1.45), Inches(1.6), 'Experience Cloud', bold=True, sub='Registration page; consent captured')
# Salesforce platform
group(s, Inches(4.2), Inches(1.65), Inches(6.35), Inches(4.65), 'SALESFORCE PLATFORM', SF_LIGHT, SF, SF)
box(s, Inches(4.4), Inches(2.05), Inches(1.9), Inches(1.15), 'Service Cloud', fill=WHITE, line=SF, bold=True, sub='Messaging + Omni-Channel routing')
box(s, Inches(6.45), Inches(2.05), Inches(1.95), Inches(1.15), 'Agentforce agent', fill=SF, line=None, color=WHITE, bold=True, sub='Understands intent, language, voice')
box(s, Inches(8.55), Inches(2.05), Inches(1.8), Inches(1.15), 'Journey engine', fill=WHITE, line=SF, bold=True, sub='Apex rules, EMI maths, EN/HI/MR')
box(s, Inches(4.4), Inches(3.4), Inches(1.9), Inches(1.15), 'RM console', fill=WHITE, line=SF, bold=True, sub='Takeover, callback queue')
box(s, Inches(6.45), Inches(3.4), Inches(1.95), Inches(1.15), 'Loan Origination', fill=WHITE, line=SF, bold=True, sub='Application, underwriting, decision, LAN')
box(s, Inches(8.55), Inches(3.4), Inches(1.8), Inches(1.15), 'Customer 360', fill=WHITE, line=SF, bold=True, sub='Contact, loans, offers')
box(s, Inches(4.4), Inches(4.75), Inches(5.95), Inches(0.55), 'Einstein Trust Layer · sharing & field security · consent · audit trail', fill=NAVY, line=None, color=WHITE, size=11, bold=True)
box(s, Inches(4.4), Inches(5.45), Inches(5.95), Inches(0.65), 'Data Cloud (next phase): unified profile, drop-off and intent segments, next best action', fill=WHITE, line=SF, size=11)
group(s, Inches(10.85), Inches(1.65), Inches(2.05), Inches(4.65), 'BANK SYSTEMS', LIGHT, LINE, GREY)
box(s, Inches(11.0), Inches(2.05), Inches(1.75), Inches(0.65), 'MuleSoft APIs', fill=NAVY, line=None, color=WHITE, bold=True, size=11)
for i, t in enumerate(['Core banking', 'CIBIL / bureau', 'NSDL · Perfios · EPFO', 'e-Sign · disbursal']):
    box(s, Inches(11.0), Inches(2.9) + i * Inches(0.82), Inches(1.75), Inches(0.65), t, size=11)
arrow(s, Inches(1.85), Inches(2.5), Inches(2.3), Inches(2.6), both=True)
arrow(s, Inches(1.85), Inches(4.8), Inches(2.3), Inches(4.65))
arrow(s, Inches(3.75), Inches(2.85), Inches(4.4), Inches(2.62), both=True, color=SF)
arrow(s, Inches(3.75), Inches(4.65), Inches(4.4), Inches(4.0), color=SF)
arrow(s, Inches(6.3), Inches(2.62), Inches(6.45), Inches(2.62), both=True, color=SF)
arrow(s, Inches(8.4), Inches(2.62), Inches(8.55), Inches(2.62), both=True, color=SF)
arrow(s, Inches(10.35), Inches(3.97), Inches(11.0), Inches(2.37), both=True, color=NAVY)
flow = ['1  Customer says "Hi" or scans the QR', '2  Agent verifies, prices the offer from the existing loan',
        '3  Accept → application, RM callback queue, LAN', '4  "Agent" → RM gets the chat with full history']
for i, f in enumerate(flow):
    box(s, Inches(0.4) + i * Inches(3.13), Inches(6.45), Inches(3.0), Inches(0.5), f, fill=WHITE, line=SF, size=11, color=NAVY, bold=True)

# ---------------------------------------------------------------- 6. Approaches + recommendation
s = prs.slides.add_slide(BLANK)
header(s, 'Two ways to run WhatsApp, and our recommendation', 'Approach and roadmap', 6)
text(s, Inches(0.6), Inches(1.3), Inches(12.2), Inches(0.35),
     'A WhatsApp number connects to one provider at a time, so the choice is one number or two.', size=13, color=GREY)
approaches = [
    ('APPROACH A', 'One number: Gupshup in front, loans on Salesforce',
     [('Customer', GREEN_LIGHT, GREEN, DARK), ('Kotak number\n(Gupshup)', AMBER_LIGHT, RGBColor(0xE0, 0x9B, 0x2D), DARK), ('Salesforce\nAgent + LOS', SF, None, WHITE)],
     ['Same verified Kotak number; nothing changes for customers', 'No number migration; alerts, OTPs and campaigns stay on Gupshup', 'Fastest to go live; this is what the LoanSeva PoC runs (BYOC)'],
     ['Gupshup builds the loan-intent hand-over and hand-back', 'Two vendors in the delivery chain; messaging billed via Gupshup'],
     'Best when: speed to market and one familiar number matter most'),
    ('APPROACH B', 'Two numbers: loans number on Salesforce native',
     [('Customer', GREEN_LIGHT, GREEN, DARK), ('Kotak Loans no.\n(Salesforce)', SF_LIGHT, SF, DARK), ('Salesforce\nAgent + LOS', SF, None, WHITE)],
     ['Fully native: one vendor for channel, AI, CRM and LOS', 'Direct with Meta, no BSP in the loan journey', 'Main number stays on Gupshup, links across with wa.me'],
     ['Two numbers: customers may be unsure which is genuine', 'Second verification and templates; history split across numbers'],
     'Best when: Kotak wants a dedicated, fully native loans number now'),
]
for i, (tag, title, chain, pros, cons, fit) in enumerate(approaches):
    x = Inches(0.6) + i * Inches(6.2)
    w = Inches(6.0)
    rec = i == 1
    box(s, x, Inches(1.75), w, Inches(3.55), '', fill=WHITE, line=SF if rec else LINE)
    text(s, x + Inches(0.2), Inches(1.85), w - Inches(0.4), Inches(0.6), [
        (tag + ('  ·  RECOMMENDED' if rec else ''), {'size': 10, 'color': RED, 'bold': True, 'after': 1}),
        (title, {'size': 15, 'color': NAVY, 'bold': True}),
    ])
    for k, (lbl, fill, line, col) in enumerate(chain):
        bx = x + Inches(0.2) + k * Inches(1.95)
        box(s, bx, Inches(2.6), Inches(1.65), Inches(0.62), lbl, fill=fill, line=line, color=col, size=10, bold=True)
        if k:
            arrow(s, bx - Inches(0.3), Inches(2.91), bx, Inches(2.91), both=True, color=SF if k == 2 else GREY)
    text(s, x + Inches(0.2), Inches(3.35), w - Inches(0.4), Inches(1.95),
         [('✓  ' + p, {'size': 11, 'color': GREEN, 'after': 2}) for p in pros] + [('✗  ' + c, {'size': 11, 'color': RGBColor(0xB0, 0x10, 0x18), 'after': 2}) for c in cons])
    box(s, x + Inches(0.2), Inches(4.7), w - Inches(0.4), Inches(0.45), fit, fill=SF_LIGHT if rec else LIGHT, line=None, color=NAVY, size=11, bold=True)
band = box(s, Inches(0.6), Inches(5.45), Inches(12.2), Inches(1.45), '', fill=NAVY, line=None)
text(s, Inches(0.8), Inches(5.52), Inches(3.0), Inches(1.3), [
    ('OUR RECOMMENDATION', {'size': 10, 'color': RGBColor(0xFF, 0xB3, 0xB6), 'bold': True, 'after': 2}),
    ('Approach B: a native loans number on Salesforce', {'size': 15, 'color': WHITE, 'bold': True}),
])
phases = [
    ('1 · Native loans number', 'Top-Up on a Kotak Loans number on Salesforce native WhatsApp; Agentforce, LOS, RM callback queue; main Gupshup bot links across'),
    ('2 · Scale the journeys', 'New loans with real bureau, KYC, Perfios and e-sign; RM live takeover; Hindi, Marathi and voice'),
    ('3 · One number', 'Move the main number to Salesforce, bring the other use cases across and retire the link; add Data Cloud next best action'),
]
for i, (k, d) in enumerate(phases):
    x = Inches(3.95) + i * Inches(2.95)
    box(s, x, Inches(5.6), Inches(2.8), Inches(1.15), '', fill=RGBColor(0x0B, 0x4A, 0x8F), line=None)
    text(s, x + Inches(0.08), Inches(5.63), Inches(2.65), Inches(1.1), [
        (k, {'size': 11, 'color': WHITE, 'bold': True, 'after': 2}),
        (d, {'size': 9.5, 'color': RGBColor(0xD6, 0xE2, 0xF0)}),
    ])

# ---------------------------------------------------------------- 7. Build effort comparison
s = prs.slides.add_slide(BLANK)
header(s, 'What Kotak has to build for conversational AI: Approach 2 is the least', 'Build effort: current state vs Approach 1 vs Approach 2', 7)
B, P, C, N = 'build', 'partner', 'config', 'na'
TONE = {B: (RED_LIGHT, RGBColor(0xB0, 0x10, 0x18)), P: (AMBER_LIGHT, RGBColor(0x9A, 0x5B, 0x00)),
        C: (GREEN_LIGHT, GREEN), N: (LIGHT, GREY)}
grid = [
    ('WhatsApp connection', (C, 'In place (Gupshup)'), (P, 'Build BYOC connector: inbound API, outbound events, auth'), (C, 'Configure: register the number in Setup')),
    ('AI / LLM layer', (B, 'Build LLM integration, prompts, guardrails, hosting'), (C, 'Configure Agentforce'), (C, 'Configure Agentforce')),
    ('Customer context for the AI', (B, 'Build APIs from CRM, LOS and core banking into the bot'), (C, 'Native: on the customer record'), (C, 'Native: on the customer record')),
    ('Journey state and LOS link', (B, 'Build a state store and the LOS integration'), (C, 'Native: journey, application, LAN'), (C, 'Native: journey, application, LAN')),
    ('Human hand-off with history', (B, 'Build context transfer to the agent inbox'), (C, 'Out of the box: Omni-Channel'), (C, 'Out of the box: Omni-Channel')),
    ('Routing between bots', (N, 'Not applicable'), (P, 'Build hand-over and hand-back with the Gupshup bot'), (C, 'Configure a wa.me link from the main bot')),
    ('Trust and governance', (B, 'Build masking, retention and audit per tool'), (P, 'Trust Layer + connector security review'), (C, 'Out of the box: Trust Layer')),
]
totals = ('To build or integrate', '5 builds · 3+ vendors', '2 partner builds + security review · 2 vendors', 'Configure only · 1 vendor for loans')
cols = [Inches(2.3), Inches(3.2), Inches(3.2), Inches(3.2)]
tbl = s.shapes.add_table(len(grid) + 3, 4, Inches(0.6), Inches(1.55), sum(cols, Emu(0)), Inches(4.7)).table
for c, w in enumerate(cols):
    tbl.columns[c].width = w


def cell_text(cell, value, size=10.5, bold=False, color=DARK, fill=None, align=PP_ALIGN.LEFT):
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.name = 'Calibri'
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


for c, (h, fill) in enumerate((('What has to exist', NAVY), ('Current state: build on Gupshup', NAVY),
                               ('Approach 1: Gupshup + Salesforce', NAVY), ('Approach 2: Salesforce native  ★', SF))):
    cell_text(tbl.cell(0, c), h, size=12, bold=True, color=WHITE, fill=fill)
for r, (label, *opts) in enumerate(grid, start=1):
    cell_text(tbl.cell(r, 0), label, bold=True, color=NAVY, fill=WHITE)
    for c, (tone, txt) in enumerate(opts, start=1):
        fill, col = TONE[tone]
        cell_text(tbl.cell(r, c), txt, color=col, fill=fill, bold=tone == B)
same = len(grid) + 1
tbl.cell(same, 1).merge(tbl.cell(same, 3))
cell_text(tbl.cell(same, 0), 'Same in every option', bold=True, color=GREY, fill=LIGHT)
cell_text(tbl.cell(same, 1), 'Loan journey design and copy; core banking, bureau, KYC and e-sign integrations (MuleSoft); RM training', color=GREY, fill=LIGHT)
last = len(grid) + 2
for c, v in enumerate(totals):
    cell_text(tbl.cell(last, c), v, size=11.5, bold=True, color=WHITE, fill=SF if c == 3 else NAVY)
legend = [(B, 'Kotak builds'), (P, 'Gupshup / partner builds'), (C, 'Configure or out of the box')]
for i, (tone, lbl) in enumerate(legend):
    fill, col = TONE[tone]
    box(s, Inches(0.6) + i * Inches(2.4), Inches(6.38), Inches(2.25), Inches(0.32), lbl, fill=fill, line=None, color=col, size=10, bold=True)
box(s, Inches(7.95), Inches(6.28), Inches(4.55), Inches(0.66), 'Recommendation: Approach 2, least to build', fill=SF, line=None, color=WHITE, size=11.5, bold=True,
    sub='One platform, no connector. Trade-off: a dedicated loans number', sub_size=9.5)

# ---------------------------------------------------------------- 8. Demo scenarios
s = prs.slides.add_slide(BLANK)
header(s, 'Three customers, one WhatsApp chat, all live in the demo', 'Demo scenarios', 8)
text(s, Inches(0.6), Inches(1.3), Inches(12.2), Inches(0.35),
     'Same agent, same number and same LOS behind all three; the customer never leaves WhatsApp.', size=13, color=GREY)
scenarios = [
    ('1 · TOP-UP', 'Priya Sharma', 'Existing personal loan · 12/12 EMIs on time · ₹5L pre-approved', NAVY,
     ['Scans a Kotak QR; LoanSeva greets her on WhatsApp',
      'Asks the rate on FAQ; switches to Hindi and back',
      'Name, mobile and OTP; sees her masked loan and track',
      'Five quick answers: job, employer, income, purpose, amount',
      'Offer ₹5L at 10.99%; edits tenure; EMI recalculates',
      'Accepts and picks an RM callback slot'],
     ['Top-Up application, pre-approved: no fresh underwriting', 'RM callback queue, Task and Event with the summary'],
     'Reference KMB-PLTU-… · RM calls at the chosen slot'),
    ('2 · NEW LOAN + DOCUMENTS', 'Rahul Deshmukh', 'Existing Kotak customer · applies for a new personal loan', SF,
     ['Picks "Apply for a new loan" after the OTP',
      'Details on file are prefilled; only missing ones asked',
      'Existing Kotak EMI is counted automatically',
      'Sees an indicative offer for the amount he asks',
      'Sends photos of PAN, bank statement, salary slips',
      'Reviews, consents and submits'],
     ['Documents land on the application, no typing needed', 'Six checks: PAN, CIBIL, Perfios, employment, FOIR, fraud', 'RM sees the stage gates and credit memo, approves'],
     'Loan account created · masked LAN sent on WhatsApp'),
    ('3 · BY VOICE', 'Asha Patil', 'New to Kotak · no pre-approved offer · prefers to speak', RED,
     ['Sends a voice note with her name and mobile',
      'Speaks the OTP, date of birth and PAN letter by letter',
      'Says email, PIN, income and amount in her own words',
      'Each reply starts with "I heard: …" so she can check',
      'Works in English, Hindi or Marathi',
      'Can still type or tap at any step'],
     ['ElevenLabs speech-to-text on each note', 'Same journey as typing; PAN masked, OTP never echoed', 'Application and underwriting as for Rahul'],
     'Full application without typing a word'),
]
for i, (tag, name, profile, colour, steps, sf, outcome) in enumerate(scenarios):
    x = Inches(0.6) + i * Inches(4.125)
    w = Inches(3.95)
    box(s, x, Inches(1.75), w, Inches(5.15), '', fill=WHITE, line=LINE)
    box(s, x, Inches(1.75), w, Inches(0.95), '', fill=colour, line=None)
    text(s, x + Inches(0.15), Inches(1.8), w - Inches(0.3), Inches(0.9), [
        (tag, {'size': 10, 'color': RGBColor(0xFF, 0xD6, 0xD8) if colour != RED else WHITE, 'bold': True, 'after': 1}),
        (name, {'size': 16, 'color': WHITE, 'bold': True, 'after': 1}),
        (profile, {'size': 9.5, 'color': WHITE}),
    ])
    text(s, x + Inches(0.15), Inches(2.82), w - Inches(0.3), Inches(0.3), 'ON WHATSAPP', size=9.5, color=GREY, bold=True)
    text(s, x + Inches(0.15), Inches(3.08), w - Inches(0.3), Inches(2.0),
         [(f'{k + 1}.  {t}', {'size': 10.5, 'after': 3}) for k, t in enumerate(steps)])
    box(s, x + Inches(0.15), Inches(5.08), w - Inches(0.3), Inches(1.05), '', fill=SF_LIGHT, line=None)
    text(s, x + Inches(0.25), Inches(5.11), w - Inches(0.5), Inches(1.0),
         [('IN SALESFORCE', {'size': 9.5, 'color': SF, 'bold': True, 'after': 2})] + [(t, {'size': 9.5, 'bullet': True, 'after': 1}) for t in sf])
    box(s, x + Inches(0.15), Inches(6.22), w - Inches(0.3), Inches(0.52), outcome, fill=GREEN_LIGHT, line=None, color=GREEN, size=10.5, bold=True)

# ---------------------------------------------------------------- 9. Backup: message flow
s = prs.slides.add_slide(BLANK)
header(s, 'How a WhatsApp message flows: Gupshup → Salesforce → Gupshup', 'Backup · Option 1a (Bring Your Own Channel)', 9)
lanes = [('Customer', 'WhatsApp', GREEN_LIGHT, GREEN, DARK), ('Meta', 'Cloud API', WHITE, LINE, DARK),
         ('Gupshup', 'BSP + BYOC connector', AMBER_LIGHT, RGBColor(0xE0, 0x9B, 0x2D), DARK),
         ('Service Cloud', 'Messaging + Omni-Channel', WHITE, SF, DARK), ('Agentforce', 'LoanSeva by Kotak', SF, None, WHITE),
         ('LoanSeva engine', 'Apex + LOS records', WHITE, SF, DARK)]
lw, lx0 = Inches(1.45), Inches(0.5)
cx = []
for i, (t, sub, fill, line, col) in enumerate(lanes):
    x = lx0 + i * Inches(1.6)
    box(s, x, Inches(1.5), lw, Inches(0.62), t, fill=fill, line=line, color=col, size=11, bold=True, sub=sub, sub_size=8)
    c = x + lw // 2
    cx.append(c)
    arrow(s, c, Inches(2.12), c, Inches(6.75), color=LINE, dashed=True, width=1)
    s.shapes[-1].line._get_or_add_ln().remove(s.shapes[-1].line._get_or_add_ln().find(qn('a:tailEnd')))
steps = [
    (0, 1, '① Text, photo or voice'),
    (1, 2, '② Webhook to the BSP'),
    (2, 3, '③ Interaction API (OAuth)'),
    (3, 4, '④ Omni-Flow routes'),
    (4, 5, '⑤ LoanSevaReplyAction'),
    (5, 4, '⑥ Reply + next step'),
    (4, 3, '⑦ Reply on session'),
    (3, 2, '⑧ Platform event (Pub/Sub)'),
    (2, 1, '⑨ Gupshup send API'),
    (1, 0, '⑩ Delivered on WhatsApp'),
]
for k, (a, b, label) in enumerate(steps):
    y = Inches(2.55) + k * Inches(0.42)
    outbound = k >= 5
    col = SF if 2 <= min(a, b) and max(a, b) <= 5 else GREY
    arrow(s, cx[a], y, cx[b], y, color=RED if outbound and col == GREY else col, width=1.75)
    lx = min(cx[a], cx[b])
    text(s, lx - Inches(0.25), y - Inches(0.3), abs(cx[b] - cx[a]) + Inches(0.5), Inches(0.26), label, size=9, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
text(s, Inches(0.5), Inches(6.75), Inches(9.5), Inches(0.3), 'Steps ①–⑤ inbound, ⑥–⑩ outbound. In the PoC a package flow on the platform event calls the demo router instead of Gupshup.', size=9, color=GREY)
group(s, Inches(10.2), Inches(1.5), Inches(2.65), Inches(5.25), 'SAME PATH, OTHER CASES', LIGHT, LINE, GREY)
side = [
    ('Proactive messages', 'Nudges and "LAN sanctioned" go out via sendConversationMessages → steps ⑧–⑩'),
    ('Human hand-off', 'Agent escalates → Omni-Channel → RM queue; RM replies in the console → steps ⑧–⑩'),
    ('Hand back to Gupshup', 'When the Messaging Session ends, Gupshup resumes its own bot'),
    ('Option 1b instead', 'Gupshup calls the Agent API directly: no session, no Omni-Channel, Gupshup sends the reply'),
    ('One call per message', 'Each customer message is one inbound API call into the same session; a 15-message chat is 15 calls'),
]
text(s, Inches(10.3), Inches(1.9), Inches(2.45), Inches(4.8),
     sum([[(t, {'size': 11, 'color': NAVY, 'bold': True, 'after': 1}), (d, {'size': 9.5, 'color': DARK, 'after': 9})] for t, d in side], []))

# ---------------------------------------------------------------- 10. Backup: native message flow
s = prs.slides.add_slide(BLANK)
header(s, 'How a WhatsApp message flows with Salesforce as the channel', 'Backup · Option 2 and end state (Salesforce native WhatsApp)', 10)
lanes = [('Customer', 'WhatsApp', GREEN_LIGHT, GREEN, DARK), ('Meta', 'Cloud API', WHITE, LINE, DARK),
         ('Service Cloud', 'WhatsApp + Omni-Channel', WHITE, SF, DARK), ('Agentforce', 'LoanSeva by Kotak', SF, None, WHITE),
         ('LoanSeva engine', 'Apex + LOS records', WHITE, SF, DARK)]
lw, lx0 = Inches(1.7), Inches(0.5)
cx = []
for i, (t, sub, fill, line, col) in enumerate(lanes):
    x = lx0 + i * Inches(1.95)
    box(s, x, Inches(1.5), lw, Inches(0.62), t, fill=fill, line=line, color=col, size=11, bold=True, sub=sub, sub_size=8)
    c = x + lw // 2
    cx.append(c)
    arrow(s, c, Inches(2.12), c, Inches(6.5), color=LINE, dashed=True, width=1)
    s.shapes[-1].line._get_or_add_ln().remove(s.shapes[-1].line._get_or_add_ln().find(qn('a:tailEnd')))
sf_band = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, lx0 + 2 * Inches(1.95) - Inches(0.08), Inches(1.42), 3 * Inches(1.95) - Inches(0.1), Inches(5.15))
sf_band.adjustments[0] = 0.03
sf_band.fill.background()
sf_band.line.color.rgb = SF
sf_band.line.width = Pt(1.25)
sf_band.line._get_or_add_ln().append(etree.fromstring('<a:prstDash xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" val="dash"/>'))
text(s, lx0 + 2 * Inches(1.95), Inches(6.6), 3 * Inches(1.95), Inches(0.3), 'Everything inside the dashed box is Salesforce', size=9, color=SF, bold=True, align=PP_ALIGN.CENTER)
steps = [
    (0, 1, '① Text, photo or voice'),
    (1, 2, '② Meta webhook to Salesforce'),
    (2, 3, '③ Omni-Flow routes'),
    (3, 4, '④ LoanSevaReplyAction'),
    (4, 3, '⑤ Reply + next step'),
    (3, 2, '⑥ Reply on session'),
    (2, 1, '⑦ Salesforce sends to Meta'),
    (1, 0, '⑧ Delivered on WhatsApp'),
]
for k, (a, b, label) in enumerate(steps):
    y = Inches(2.6) + k * Inches(0.5)
    col = SF if min(a, b) >= 2 else (GREY if k < 4 else RED)
    arrow(s, cx[a], y, cx[b], y, color=col, width=1.75)
    lx = min(cx[a], cx[b])
    w = Inches(0.13) * len(label)
    tb = text(s, (cx[a] + cx[b]) // 2 - w // 2, y - Inches(0.32), w, Inches(0.26), label, size=10, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    tb.fill.solid()
    tb.fill.fore_color.rgb = WHITE
group(s, Inches(10.4), Inches(1.5), Inches(2.45), Inches(2.75), 'VS OPTION 1A', SF_LIGHT, SF, SF)
text(s, Inches(10.5), Inches(1.88), Inches(2.3), Inches(2.35), [(t, {'size': 9.5, 'color': DARK, 'after': 5, 'bullet': True}) for t in (
    'No BSP and no connector: two fewer hops',
    'No inbound API calls into the org and no outbound platform events',
    'Number registered with Meta from Salesforce Setup; templates managed in Salesforce',
    'Usage via Salesforce messaging packs plus Meta\'s fees',
)])
group(s, Inches(10.4), Inches(4.4), Inches(2.45), Inches(2.35), 'SAME AS OPTION 1A', LIGHT, LINE, GREY)
text(s, Inches(10.5), Inches(4.78), Inches(2.3), Inches(1.95), [(t, {'size': 9.5, 'color': DARK, 'after': 5, 'bullet': True}) for t in (
    'Agentforce and the engine run once per customer message',
    'RM hand-off via Omni-Channel in the console',
    'Nudges via sendConversationMessages',
    'Meta\'s messaging tiers and throughput limits',
)])

out = 'docs/LoanSeva_WhatsApp_Banking_Why_Salesforce.pptx'
prs.save(out)
print(out)
