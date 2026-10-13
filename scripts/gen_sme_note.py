"""Builds docs/LoanSeva_WhatsApp_Channel_Evaluation_Note.docx: a note for a Salesforce SME comparing Gupshup as the
WhatsApp channel (BYOC or Agent API) with Salesforce native WhatsApp. Run: python3 scripts/gen_sme_note.py"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x00, 0x38, 0x74)
RED = RGBColor(0xED, 0x1C, 0x24)
GREY = RGBColor(0x5F, 0x63, 0x68)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2)
base = doc.styles['Normal']
base.font.name = 'Calibri'
base.font.size = Pt(10.5)
for name, size, color in (('Heading 1', 16, NAVY), ('Heading 2', 12.5, NAVY), ('Heading 3', 11, RED)):
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = True


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear')
    sh.set(qn('w:color'), 'auto')
    sh.set(qn('w:fill'), hex_fill)
    tcPr.append(sh)


def table(rows, widths=None, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ''
            run = cell.paragraphs[0].add_run(val)
            run.font.size = Pt(9.5)
            if header and r == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(cell, '003874')
            elif c == 0:
                run.font.bold = True
            if widths:
                cell.width = Cm(widths[c])
    doc.add_paragraph()
    return t


def para(txt, bold_lead=None, style=None):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(txt)
    return p


def bullets(items, style='List Bullet'):
    for i in items:
        if isinstance(i, tuple):
            para(i[1], bold_lead=i[0] + ' ', style=style)
        else:
            doc.add_paragraph(i, style=style)


def note(txt):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run('Note: ')
    r.bold = True
    r.font.color.rgb = RED
    p.add_run(txt).italic = True


# ------------------------------------------------------------------ Cover
p = doc.add_paragraph()
r = p.add_run('KOTAK MAHINDRA BANK × SALESFORCE · FOR SME REVIEW')
r.bold = True
r.font.color.rgb = RED
r.font.size = Pt(11)
p = doc.add_paragraph()
r = p.add_run('WhatsApp channel for LoanSeva: Gupshup or Salesforce native?')
r.bold = True
r.font.size = Pt(22)
r.font.color.rgb = NAVY
p = doc.add_paragraph()
r = p.add_run('An evaluation note on architecture, API and event consumption, cost drivers and customer benefit')
r.font.size = Pt(13)
r.font.color.rgb = GREY
table([
    ('Purpose', 'Ask a Salesforce Digital Engagement / Agentforce SME to confirm or correct the analysis below and recommend the channel for Kotak\'s loan journeys.'),
    ('Context', 'LoanSeva proof of concept: WhatsApp Top-Up and new-loan journeys with Agentforce, an Apex journey engine and an LOS, in org kotak-demo (00DgL00000eqlv3UAA).'),
    ('Our view', 'Salesforce native WhatsApp for the loan journey (Option 2). Section 7 lists what must be confirmed before we commit.'),
    ('Status of facts', 'Limit-counting and licensing points are our understanding, not verified; they are listed as questions in Section 7. The PoC runs on BYOC; native WhatsApp has not been tested in this org.'),
], widths=[3.2, 13.3], header=False)

# ------------------------------------------------------------------ 1
doc.add_heading('1. The decision', level=1)
para('Kotak runs WhatsApp today through Gupshup, its Business Solution Provider (BSP). The LoanSeva PoC shows the loan journey working on '
     'Salesforce: Agentforce holds the conversation, an Apex engine owns the journey and the LOS records the application. The question is '
     'how WhatsApp messages should reach Salesforce in production. A WhatsApp number can be connected to only one provider at a time, so the '
     'options are:')
table([
    ('Option', 'How it works', 'Number'),
    ('1a · Gupshup + BYOC', 'Gupshup stays the BSP. Loan conversations are forwarded to Salesforce through Bring Your Own Channel (BYOC). '
     'Omni-Channel routes them to the agent or an RM in the Service Console.', 'Existing Kotak number'),
    ('1b · Gupshup + Agent API', 'Gupshup stays the BSP and calls the Agentforce Agent API directly for each message. No Salesforce messaging channel, '
     'no Omni-Channel and no console takeover; hand-off stays in Gupshup\'s inbox.', 'Existing Kotak number'),
    ('2 · Salesforce native', 'A Kotak Loans WhatsApp number is registered on Salesforce Messaging for WhatsApp. Meta talks to Salesforce directly; '
     'Gupshup is not in the loan journey.', 'Dedicated loans number (main number can move later)'),
], widths=[3.4, 10.0, 3.1])

# ------------------------------------------------------------------ 2
doc.add_heading('2. What the PoC runs today (evidence)', level=1)
para('The PoC uses option 1a with a partner BYOC package, which stands in for what Gupshup would build:')
bullets([
    ('Channel.', 'Messaging channel WhatsApp_BYOC on ConversationChannelDefinition SEs_WhatsAppBYOC (managed package wpbyocdemo3). Consent and routing are owned by Salesforce.'),
    ('Inbound.', 'Every customer message is posted by the external router to the Salesforce Interaction Service API, authenticated with OAuth (JWT bearer).'),
    ('Outbound.', 'Every agent or RM message is published as the platform event wpbyocdemo3__WP_BYOC_Meta_Event__e, relayed by the package flow '
     '"[WA Meta] Send Outbound Event" to the router, which calls Meta. The flow is protected in the managed package, so we could not inspect how the router receives it.'),
    ('Routing and agent.', 'Omni-Flow LoanSeva_WhatsApp_Inbound_Routing hands the session to the LoanSeva_by_Kotak agent; each customer reply runs one '
     'agent action (LoanSevaReplyAction → LoanSevaEngine.reply()).'),
    ('Proactive messages.', 'Nudges and offers use sendConversationMessages with the LoanSeva_Text messaging component, then go out through the same platform event.'),
    ('Rich messages.', 'The journey uses numbered replies. Whether tappable WhatsApp buttons and lists render over BYOC depends on the router; this was not proven in the PoC.'),
])
para('Allocations in the PoC org at the time of writing (Kotak\'s production org will differ):')
table([
    ('Limit (sf org list limits)', 'PoC org allocation'),
    ('DailyApiRequests (24-hour REST/SOAP allocation)', '441,450'),
    ('DailyDeliveredPlatformEvents (to external subscribers)', '75,000'),
    ('HourlyPublishedPlatformEvents', '450,000'),
    ('MonthlyPlatformEventsUsageEntitlement', '1,950,000'),
], widths=[10.5, 6.0])

# ------------------------------------------------------------------ 3
doc.add_heading('3. Message path in each option', level=1)
doc.add_heading('Option 1a · Gupshup + BYOC', level=3)
bullets([
    'Customer → Meta → Gupshup webhook. Gupshup decides it is a loan conversation (its hand-over logic).',
    'Gupshup → Salesforce Interaction Service API: one inbound call per customer message, plus calls for attachments and, if sent, delivery and read receipts.',
    'Salesforce creates or reuses the MessagingSession, Omni-Flow routes it, Agentforce runs one action per message.',
    'Agent reply → platform event → Gupshup (subscriber) → Meta → customer. One event per outbound message.',
    'RM takeover in the Service Console works as in the PoC. Hand-back to the main Gupshup bot is Gupshup\'s logic.',
])
doc.add_heading('Option 1b · Gupshup + Agent API', level=3)
bullets([
    'Customer → Meta → Gupshup. Gupshup starts an Agent API session and posts each customer message to it (one call per message, plus session start and end).',
    'Agentforce runs one action per message and returns the reply in the API response; Gupshup sends it to Meta.',
    'No MessagingSession, no Omni-Channel routing, no console takeover and no Salesforce transcript unless Gupshup writes one back (more API calls).',
    'Proactive nudges from Salesforce have no Salesforce channel to go out on; they need a Gupshup API call (an outbound callout from Salesforce).',
])
doc.add_heading('Option 2 · Salesforce native WhatsApp', level=3)
bullets([
    'Customer → Meta → Salesforce Messaging (Meta\'s webhook is received by Salesforce\'s messaging platform, not by Kotak\'s org API).',
    'MessagingSession, Omni-Flow and Agentforce exactly as in 1a. The LoanSeva engine, actions and LOS are unchanged.',
    'Reply → Salesforce Messaging → Meta → customer. No platform event, no partner, no connector.',
    'Templates, opt-in and the number are managed in Salesforce Setup. Nudges use the same sendConversationMessages call the PoC already uses.',
])

# ------------------------------------------------------------------ 4
doc.add_heading('4. API and event consumption', level=1)
para('This is the main concern with keeping Gupshup in front. In options 1a and 1b, every message in every conversation crosses into the org '
     'through an API, so consumption grows with chat volume, not with the number of customers or users. In option 2 the same traffic stays inside '
     'Salesforce\'s messaging platform.')
doc.add_heading('4.1 Per conversation', level=2)
para('A Top-Up conversation in the PoC is about 12 customer replies and about 15 outbound messages. A new-loan conversation is longer (about 20 replies '
     'plus 3 document photos). Using the Top-Up as the example:')
table([
    ('Per conversation', '1a · Gupshup + BYOC', '1b · Gupshup + Agent API', '2 · Native'),
    ('Inbound customer messages', '~12 Interaction Service API calls', '~12 Agent API calls + session start/end', '0 org API calls'),
    ('Delivery / read receipts', 'Up to ~30 calls, if Gupshup forwards them', 'Not tracked in Salesforce', '0 (handled by Salesforce Messaging)'),
    ('Attachments (new loan)', 'Extra calls per photo or voice note', 'Depends on Gupshup\'s design', '0'),
    ('Outbound messages', '~15 platform events published and delivered', 'Returned in the API response', '0 platform events'),
    ('Transcript in Salesforce', 'Yes (MessagingSession)', 'Only if Gupshup writes it back (more calls)', 'Yes (MessagingSession)'),
    ('Agentforce actions', '~12', '~12', '~12'),
], widths=[3.6, 4.3, 4.3, 4.3])
note('Agentforce runs once per customer message in every option, so Agentforce consumption (Flex Credits) does not change with the channel. '
     'What changes is the API and event traffic around it.')
doc.add_heading('4.2 At Kotak scale (illustrative)', level=2)
para('Assuming ~42 inbound calls per conversation in 1a (12 messages + 30 receipts) and ~15 outbound events:')
table([
    ('Loan conversations per day', '1a inbound calls / day', '1a events delivered / day', '2 · native'),
    ('1,000', '~42,000', '~15,000', '0 / 0'),
    ('5,000', '~210,000', '~75,000 (the PoC org\'s whole daily delivery allocation)', '0 / 0'),
    ('20,000 (campaign peak)', '~840,000', '~300,000', '0 / 0'),
], widths=[4.0, 3.8, 5.6, 3.1])
para('What this means for Kotak if these calls count toward the org\'s allocations:')
bullets([
    ('Shared limit.', 'The 24-hour API allocation is shared with every other integration in Kotak\'s org (core banking, MuleSoft, data loads, mobile). WhatsApp traffic spikes after a campaign send could starve them, or be throttled itself.'),
    ('Paid add-ons.', 'Higher API or platform-event delivery allocations are bought as add-ons, so the cost grows with chat volume.'),
    ('Monitoring.', 'Kotak would need limit alerts and back-off and retry in Gupshup\'s connector; a throttled call is a lost or delayed customer message.'),
    ('Option 2.', 'None of the above for the loan traffic: usage is billed through Salesforce messaging packs, which are sized and priced for message volume.'),
])
note('Whether Interaction Service API calls (BYOC inbound) and Agent API calls count toward DailyApiRequests is the most important question for the SME (Section 7, Q1-Q2). '
     'If they do not, the cost argument in 4.2 is weaker, but the event-delivery and moving-parts points still apply.')

# ------------------------------------------------------------------ 5
doc.add_heading('5. Side-by-side evaluation', level=1)
table([
    ('Criterion', '1a · Gupshup + BYOC', '1b · Gupshup + Agent API', '2 · Salesforce native'),
    ('Org API / event consumption', 'Per message, grows with volume', 'Per message, grows with volume', 'None for messaging'),
    ('What Kotak or partners build', 'Gupshup: BYOC connector, loan-intent hand-over and hand-back; security review', 'Gupshup: Agent API client, session handling, hand-off, transcript write-back', 'Configuration only: channel, number, templates, routing'),
    ('Vendors in the loan conversation', 'Meta, Gupshup, Salesforce', 'Meta, Gupshup, Salesforce', 'Meta, Salesforce'),
    ('RM live takeover with full history', 'Yes, Service Console', 'No, Gupshup inbox', 'Yes, Service Console'),
    ('Transcript on the customer record', 'Yes', 'Only if written back', 'Yes'),
    ('Proactive nudges and offers', 'Yes, via sendConversationMessages', 'Needs Gupshup API callouts', 'Yes, via sendConversationMessages'),
    ('Rich WhatsApp messages (buttons, lists)', 'Depends on Gupshup\'s connector', 'Depends on Gupshup', 'Native messaging components (confirm the supported types)'),
    ('Templates, opt-in, consent', 'Split between Gupshup and Salesforce', 'In Gupshup', 'In Salesforce'),
    ('Customer data in transit', 'Passes through Gupshup', 'Passes through Gupshup', 'Meta to Salesforce only'),
    ('Fault points per message', 'Meta, Gupshup, API, event, router', 'Meta, Gupshup, API', 'Meta, Salesforce'),
    ('Number', 'Same Kotak number', 'Same Kotak number', 'Dedicated loans number (or migrate the main one)'),
    ('Time to first live journey', 'Depends on Gupshup building the connector', 'Depends on Gupshup', 'Number verification and template approval'),
    ('Commercials', 'Salesforce + Gupshup markup + Meta', 'Salesforce + Gupshup + Meta', 'Salesforce messaging packs + Meta'),
], widths=[3.7, 4.3, 4.3, 4.2])

# ------------------------------------------------------------------ 6
doc.add_heading('6. Why native is better for Kotak and its customers', level=1)
doc.add_heading('For Kotak', level=3)
bullets([
    ('Predictable cost.', 'Message volume is billed through messaging packs, not through API and event allocations shared with the rest of the org. A campaign that triples chat volume does not risk the core banking integrations.'),
    ('Least to build.', 'No connector, no hand-over or hand-back logic and no partner security review. The PoC\'s agent, engine and LOS move across unchanged; only the channel changes.'),
    ('One accountable vendor.', 'One support path from Meta to the loan record. Fewer hops, fewer places a message can be lost and one set of logs.'),
    ('Data protection.', 'Loan conversations carry PAN, income and bank statements. With native WhatsApp they go from Meta straight to Salesforce, which simplifies DPDP and RBI outsourcing reviews (to be confirmed by Kotak\'s risk team).'),
    ('One view of the customer.', 'Transcript, consent, journey, application and RM work on the same record, ready for Data Cloud next best action later.'),
])
doc.add_heading('For the customer', level=3)
bullets([
    ('Faster replies.', 'One hop fewer each way (no partner and no event relay), so less latency and fewer timeouts at busy times.'),
    ('Seamless hand-off.', 'When they ask for a person, the RM picks up in the same chat with the full history; the customer doesn\'t repeat themselves.'),
    ('Richer messages.', 'Buttons, lists, documents and voice notes handled natively, not dependent on what a connector supports.'),
    ('Fewer failures.', 'No dropped or delayed messages because an API allocation ran out during a campaign.'),
])
doc.add_heading('Trade-offs to weigh', level=3)
bullets([
    ('Two numbers.', 'A dedicated loans number means customers see two Kotak numbers until the main number moves. Mitigation: a green-tick verified display name, and a wa.me link from the main Gupshup bot.'),
    ('New number ramp-up.', 'Meta\'s messaging tiers start low for a new number, so campaign volume has to build up over weeks.'),
    ('Gupshup features.', 'Other bots, campaigns and analytics on the main number stay on Gupshup until a later phase.'),
    ('Commercials.', 'Salesforce messaging packs replace Gupshup\'s markup on loan traffic; the net cost needs the actual quotes (see the cost model).'),
])

# ------------------------------------------------------------------ 7
doc.add_heading('7. Questions for the SME', level=1)
bullets([
    ('Q1 · BYOC inbound.', 'Do Interaction Service API calls from a BYOC partner count toward the org\'s DailyApiRequests allocation? Are there separate rate limits on that endpoint?'),
    ('Q2 · Agent API.', 'Do Agentforce Agent API calls (session start, message, end) count toward DailyApiRequests, and what are the per-org rate limits?'),
    ('Q3 · Outbound events.', 'For a BYOC channel defined with a custom platform event, does each outbound message count toward event publishing and DailyDeliveredPlatformEvents when the partner subscribes over the Pub/Sub API?'),
    ('Q4 · Receipts and attachments.', 'Do delivery and read receipts and inbound attachments each need a separate Interaction Service API call?'),
    ('Q5 · Native limits.', 'Confirm that native Messaging for WhatsApp uses no org API calls or customer-visible platform events for inbound and outbound messages. Are there any org-level throughput limits beyond Meta\'s tiers?'),
    ('Q6 · Licensing.', 'Which SKUs are needed for 1a vs 2 (Digital Engagement / messaging add-on, messaging packs, BYOC entitlement)? How are messaging packs metered: per conversation or per message, and are Meta fees passed through?'),
    ('Q7 · Rich messages.', 'Which WhatsApp interactive types (reply buttons, list messages, templates with buttons, documents, audio) are supported natively and over BYOC today?'),
    ('Q8 · Agentforce on both channels.', 'Any functional difference for an Agentforce Service Agent on native WhatsApp vs BYOC (attachments, voice notes, escalation, session timeouts)?'),
    ('Q9 · Number migration.', 'What is the process and downtime to move an existing verified number from Gupshup to Salesforce later, keeping the display name, green tick and templates?'),
    ('Q10 · Data residency.', 'Where are WhatsApp messages processed and stored for an India org on native WhatsApp (Hyperforce India)?'),
    ('Q11 · Gupshup capability.', 'Is there an existing Gupshup BYOC connector for Salesforce, or would it be a new build?'),
])

# ------------------------------------------------------------------ 8
doc.add_heading('8. Our recommendation, subject to the answers', level=1)
para('Run the loan journey on Salesforce native WhatsApp on a dedicated Kotak Loans number (Option 2). Keep the main number on Gupshup, linked across with wa.me, '
     'and move it to Salesforce in a later phase. It removes per-message API and event consumption from the org, needs no partner build, keeps one vendor '
     'between Meta and the loan record, and gives customers faster replies and a seamless RM hand-off.')
para('Option 1a stays the fallback if Kotak must keep a single number now. In that case, size API and platform-event allocations from Section 4 before go-live, '
     'and require Gupshup\'s connector to retry with back-off and not forward receipts it doesn\'t need. Option 1b is not recommended for loans: it loses console '
     'takeover, the Salesforce transcript and Salesforce-initiated nudges.')
note('Related documents: LoanSeva_WhatsApp_Banking_Why_Salesforce.pptx (slides 6, 7, 9 and 10), LoanSeva_WhatsApp_Cost_Model.xlsx and LoanSeva_Build_and_Setup_Guide.docx. '
     'Per-conversation counts come from the PoC journeys; volumes in 4.2 are illustrative.')

out = 'docs/LoanSeva_WhatsApp_Channel_Evaluation_Note.docx'
doc.save(out)
print(out)
